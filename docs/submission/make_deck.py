"""Six-slide CivicOps deck (16:9). Claims limited to what the synthetic build demonstrates."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os
M=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','media')+'/'
INK=RGBColor(0x0a,0x1e,0x2a); TEAL=RGBColor(0x07,0x5d,0x61); MUTED=RGBColor(0x52,0x69,0x77)
p=Presentation(); p.slide_width=Inches(13.333); p.slide_height=Inches(7.5)
def slide(title, bullets, image=None, kicker=''):
    s=p.slides.add_slide(p.slide_layouts[6])
    if kicker:
        t=s.shapes.add_textbox(Inches(.6),Inches(.4),Inches(8),Inches(.4)).text_frame; t.text=kicker
        r=t.paragraphs[0].runs[0]; r.font.size=Pt(13); r.font.bold=True; r.font.color.rgb=TEAL
    t=s.shapes.add_textbox(Inches(.6),Inches(.8),Inches(6.2 if image else 12),Inches(1.2)).text_frame; t.word_wrap=True; t.text=title
    r=t.paragraphs[0].runs[0]; r.font.size=Pt(34); r.font.bold=True; r.font.color.rgb=INK
    b=s.shapes.add_textbox(Inches(.6),Inches(2.2),Inches(6.0 if image else 12),Inches(4.8)).text_frame; b.word_wrap=True
    for i,x in enumerate(bullets):
        para=b.paragraphs[0] if i==0 else b.add_paragraph(); para.text='• '+x; para.space_after=Pt(10)
        para.runs[0].font.size=Pt(18); para.runs[0].font.color.rgb=MUTED
    if image: s.shapes.add_picture(image,Inches(7.0),Inches(1.1),width=Inches(5.8))
    f=s.shapes.add_textbox(Inches(.6),Inches(6.9),Inches(12),Inches(.4)).text_frame; f.text='Culltron CivicOps · synthetic demo data · simulated providers'
    f.paragraphs[0].runs[0].font.size=Pt(10); f.paragraphs[0].runs[0].font.color.rgb=MUTED
s=p.slides.add_slide(p.slide_layouts[6]); s.shapes.add_picture('cover.png',0,0,width=p.slide_width,height=p.slide_height)
slide('Reports get lost between call, ticket and crew',
 ['Residents repeat themselves across phone, WhatsApp and forms','Dispatchers cannot see why a case was routed or who decided','Agencies (utilities, Metro Police, Traffic, Fire, EMS, Disaster) work in separate systems','AI that dispatches on its own is not acceptable for public safety'],M+'resident-desk.png','THE PROBLEM')
slide('One front door, one case, one reviewed decision',
 ['Voice, SMS, WhatsApp, web and Open311 create one case with one reference','Bounded agents extract, triage and recommend; every run records provider, confidence and rationale','A human approves the response and assigns an eligible crew','Live map, crew PWA and Android app with consent-based location'],M+'control-room.png','THE PRODUCT')
slide('Agencies, SOS and cameras, with a human at every step',
 ['Cross-agency incidents: one commander, one timeline, human-decided mutual aid','Metro Police details redacted outside the agency','SOS alerts the control room; never replaces 10111, 10177 or 112','Camera adapter accepts scene events only: no faces, plates or tracking; a human confirms each event'],M+'agency-board.png','EVENT-WEEK WORK')
slide('Architecture and honest status',
 ['ElevenLabs channels → .NET API with bounded agents → Zendesk outbox → human control room → verified audit','Tenancy, SSO/RBAC, POPIA controls, PostgreSQL, backup/restore, OpenTelemetry','1,633 automated tests passing plus browser and API smoke checks (13 Oct 2026)','Not yet live: AMD inference (no endpoint for this build), ElevenLabs, Zendesk, SMS/WhatsApp providers, cameras, deployment'],M+'audit-verify.png','HOW IT WORKS')
slide('Next: a private pilot on our own desk',
 ['Culltron\'s own support desk is the first intended live proof','Municipal onboarding: identity mapping, real workforce, retention and POPIA agreements','Per-agency training packs feeding the VR training layer','We report measured results only after a pilot; no impact numbers are claimed today'],None,'WHAT\'S NEXT')
p.save('civicops-deck.pptx')
